import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.shared import Inches

from hollyocr.core.extraction import extract_docx_text


class DocxExtractionTests(unittest.TestCase):
    def extract(self, document):
        with tempfile.TemporaryDirectory(prefix="hollyocr_docx_tests_") as directory:
            path = Path(directory) / "documento.docx"
            document.save(path)
            return extract_docx_text(path)

    def test_body_keeps_paragraphs_and_tables_in_document_order(self):
        document = Document()
        document.add_paragraph("Parágrafo antes da primeira tabela.")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Primeira tabela."
        document.add_paragraph("Parágrafo entre as tabelas.")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Segunda tabela."
        document.add_paragraph("Parágrafo depois da segunda tabela.")
        self.assertEqual(
            self.extract(document),
            "Parágrafo antes da primeira tabela.\nPrimeira tabela.\n"
            "Parágrafo entre as tabelas.\nSegunda tabela.\n"
            "Parágrafo depois da segunda tabela.",
        )

    def test_nested_table_preserves_content_and_cell_block_order(self):
        document = Document()
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "Antes da tabela aninhada."
        nested = cell.add_table(rows=1, cols=2)
        nested.cell(0, 0).text = "Produto exclusivo"
        nested.cell(0, 1).text = "R$ 1900,00"
        cell.add_paragraph("Depois da tabela aninhada.")
        self.assertEqual(
            self.extract(document),
            "Antes da tabela aninhada.\nProduto exclusivo\tR$ 1900,00\n"
            "Depois da tabela aninhada.",
        )

    def test_header_footer_tables_and_enabled_variants_are_extracted(self):
        document = Document()
        section = document.sections[0]
        section.different_first_page_header_footer = True
        document.settings.odd_and_even_pages_header_footer = True
        containers = (
            (section.header, "Tabela do cabeçalho padrão"),
            (section.footer, "Tabela do rodapé padrão"),
            (section.first_page_header, "Cabeçalho da primeira página"),
            (section.first_page_footer, "Rodapé da primeira página"),
            (section.even_page_header, "Cabeçalho das páginas pares"),
            (section.even_page_footer, "Rodapé das páginas pares"),
        )
        for container, text in containers:
            container.add_table(rows=1, cols=1, width=Inches(4)).cell(0, 0).text = text
        document.add_paragraph("Conteúdo principal.")
        text = self.extract(document)
        for _, expected in containers:
            self.assertEqual(text.count(expected), 1)
        self.assertLess(text.index("Tabela do cabeçalho padrão"), text.index("Conteúdo principal."))
        self.assertLess(text.index("Conteúdo principal."), text.index("Tabela do rodapé padrão"))

    def test_linked_headers_and_footers_are_emitted_once(self):
        document = Document()
        first = document.sections[0]
        first.header.paragraphs[0].text = "Cabeçalho compartilhado entre seções."
        first.footer.paragraphs[0].text = "Rodapé compartilhado entre seções."
        document.add_paragraph("Primeira seção.")
        second = document.add_section(WD_SECTION_START.NEW_PAGE)
        self.assertTrue(second.header.is_linked_to_previous)
        self.assertTrue(second.footer.is_linked_to_previous)
        document.add_paragraph("Segunda seção.")
        text = self.extract(document)
        self.assertEqual(text.count("Cabeçalho compartilhado entre seções."), 1)
        self.assertEqual(text.count("Rodapé compartilhado entre seções."), 1)
        self.assertIn("Primeira seção.", text)
        self.assertIn("Segunda seção.", text)

    def test_merged_cells_do_not_duplicate_text_or_nested_tables(self):
        document = Document()
        table = document.add_table(rows=3, cols=3)
        horizontal = table.cell(0, 0).merge(table.cell(0, 1))
        horizontal.text = "Célula horizontal exclusiva."
        horizontal.add_table(rows=1, cols=1).cell(0, 0).text = "Tabela aninhada exclusiva."
        vertical = table.cell(1, 0).merge(table.cell(2, 0))
        vertical.text = "Célula vertical exclusiva."
        table.cell(0, 2).text = "Última coluna."
        table.cell(1, 1).text = "Segunda linha."
        table.cell(2, 1).text = "Terceira linha."
        text = self.extract(document)
        for expected in (
            "Célula horizontal exclusiva.",
            "Tabela aninhada exclusiva.",
            "Célula vertical exclusiva.",
        ):
            self.assertEqual(text.count(expected), 1)
        self.assertIn("Última coluna.", text)
        self.assertIn("Segunda linha.", text)
        self.assertIn("Terceira linha.", text)
        self.assertIn("\n\tTerceira linha.\t", text)

    def test_empty_edge_cells_preserve_column_positions(self):
        document = Document()
        document.add_paragraph("Antes da tabela.")
        table = document.add_table(rows=1, cols=3)
        table.cell(0, 1).text = "R$ 100,00"
        document.add_paragraph("Depois da tabela.")
        self.assertEqual(
            self.extract(document),
            "Antes da tabela.\n\tR$ 100,00\t\nDepois da tabela.",
        )


if __name__ == "__main__":
    unittest.main()
